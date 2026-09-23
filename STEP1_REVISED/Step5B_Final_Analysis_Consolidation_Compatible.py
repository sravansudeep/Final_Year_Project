from pathlib import Path
import math
from statistics import mean, stdev

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent

STEP1_FILE = BASE_DIR / "Step1_Stratified_Split.xlsx"
STEP3_FILE = BASE_DIR / "Step3_RPMOnly_Interpolation_Validation.xlsx"
STEP4_FILE = BASE_DIR / "Step4_RPMOnly_Extrapolation_Validation.xlsx"
STEP5_FILE = BASE_DIR / "Step5_Final_Experimental_Holdout_Validation.xlsx"

OUTPUT_FILE = BASE_DIR / "Step5B_Final_Analysis_Consolidated.xlsx"

TARGETS = ["hc_hi", "td_to", "Vch"]
DIAMETERS = [6, 8, 10, 12, 14, 16]
GRID_RPM = list(range(0, 261, 5))


HEADER_FILL = "1F4E78"
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FILL = "17365D"
TITLE_FONT = Font(bold=True, color="FFFFFF", size=14)
THIN = Side(style="thin", color="D9E2F3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def read_xlsx(path, sheet_name):
    return pd.read_excel(path, sheet_name=sheet_name)


def clean_number(value):
    if pd.isna(value):
        return None
    try:
        f = float(value)
        if f.is_integer():
            return int(f)
        return f
    except Exception:
        return value


def style_sheet(ws, freeze=True, max_width=32):
    ws.freeze_panes = "A2" if freeze else None

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cell.border = BORDER

    for row in ws.iter_rows():
        for cell in row:
            cell.border = BORDER
            if cell.row > 1:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in col_cells[:100]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), max_width)


def add_excel_table(ws, name):
    if ws.max_row < 2 or ws.max_column < 1:
        return

    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    tab = Table(displayName=name, ref=ref)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    tab.tableStyleInfo = style
    ws.add_table(tab)


def write_df(wb, sheet_name, df, table=True):
    ws = wb.create_sheet(sheet_name)

    if df.empty:
        ws["A1"] = "No records"
        return ws

    for col_idx, col in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=str(col))
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cell.border = BORDER

    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        for col_idx, value in enumerate(row, 1):
            v = clean_number(value)
            cell = ws.cell(row=row_idx, column=col_idx, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    ws.freeze_panes = "A2"

    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in list(col_cells)[:100]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[letter].width = min(
            max(max_len + 2, 10),
            32
        )

    if table:
        add_excel_table(
            ws,
            "T_" + "".join(
                c if c.isalnum() else "_"
                for c in sheet_name
            )[:25]
        )

    return ws


def make_experimental(train, holdout):
    train_part = train[
        ["Point_ID", "d", "Rpm"] + TARGETS
    ].copy()
    train_part["Split"] = "Train_Anchors"
    train_part["Data_Status"] = "MEASURED"
    train_part["ML_Default_Use"] = "YES"

    holdout_part = holdout[
        ["Point_ID", "d", "Rpm"] + TARGETS
    ].copy()
    holdout_part["Split"] = "Test_Holdout"
    holdout_part["Data_Status"] = "MEASURED"
    holdout_part["ML_Default_Use"] = "NO"

    out = pd.concat(
        [train_part, holdout_part],
        ignore_index=True
    ).sort_values("Point_ID")

    return out.reset_index(drop=True)


def make_rpm_interpolation(step3_grid, train, holdout):
    train_lookup = {
        (float(r["d"]), float(r["Rpm"])): r
        for _, r in train.iterrows()
    }

    holdout_coords = {
        (float(r["d"]), float(r["Rpm"]))
        for _, r in holdout.iterrows()
    }

    # Step 3 grid already contains target/model/source information.
    # Keep one row per coordinate and pivot targets.
    grid = step3_grid.copy()

    if "Target" in grid.columns:
        rows = []

        for (d, rpm), group in grid.groupby(["d", "Rpm"]):
            if (float(d), float(rpm)) in holdout_coords:
                continue

            row = {
                "d": d,
                "Rpm": rpm,
                "hc_hi": None,
                "td_to": None,
                "Vch": None,
                "hc_hi_Source": None,
                "td_to_Source": None,
                "Vch_Source": None,
                "hc_hi_Model": None,
                "td_to_Model": None,
                "Vch_Model": None,
            }

            for _, g in group.iterrows():
                target = str(g["Target"])

                if target not in TARGETS:
                    continue

                status = str(g.get("Status", ""))
                value = g.get("Predicted")

                if status == "MEASURED_TRAIN_POINT":
                    lookup = train_lookup.get(
                        (float(d), float(rpm))
                    )
                    if lookup is not None:
                        value = lookup[target]
                    row[f"{target}_Source"] = "MEASURED_TRAIN"
                    row[f"{target}_Model"] = "EXPERIMENTAL"
                else:
                    row[f"{target}_Source"] = (
                        "RPM_INTERPOLATION"
                        if pd.notna(value)
                        else "UNAVAILABLE"
                    )
                    row[f"{target}_Model"] = g.get("Method")

                row[target] = value

            row["ML_Default_Use"] = "NO"
            row["Graph_Use"] = "YES"
            rows.append(row)

        return pd.DataFrame(rows).sort_values(
            ["d", "Rpm"]
        ).reset_index(drop=True)

    raise ValueError(
        "Step3 RPMOnly_Grid does not contain the expected Target column."
    )


def make_extrapolation(step4_extrap):
    cols = [
        "Target",
        "Method",
        "d",
        "Rpm",
        "Predicted",
        "Physical_Violation",
        "Applicability",
        "Status",
    ]

    out = step4_extrap.copy()

    keep = [c for c in cols if c in out.columns]
    out = out[keep].copy()

    out["ML_Default_Use"] = "NO"
    out["Graph_Use"] = "YES"

    return out


def make_holdout_check(step5_predictions):
    cols = [
        "Target",
        "Method",
        "Point_ID",
        "d",
        "Rpm",
        "Actual",
        "Predicted",
        "Residual",
        "Prediction_Valid",
        "Physical_Violation",
    ]

    keep = [c for c in cols if c in step5_predictions.columns]
    return step5_predictions[keep].copy()


def make_model_history(step2_summary, step3_summary,
                       step4_one, step4_multi, step5_summary):
    rows = []

    for _, r in step2_summary.iterrows():
        rows.append({
            "Stage": "Step 2 LOOCV",
            "Target": r.get("Target"),
            "Method": r.get("Method"),
            "Coverage_percent": r.get("Coverage_percent"),
            "R2": r.get("R2"),
            "RMSE": r.get("RMSE"),
            "MAE": r.get("MAE"),
            "MAPE_percent": r.get("MAPE_percent"),
            "Physical_Violations": None,
            "Interpretation": "Train-only method selection"
        })

    for _, r in step3_summary.iterrows():
        rows.append({
            "Stage": "Step 3 RPM interpolation",
            "Target": r.get("Target"),
            "Method": r.get("Method"),
            "Coverage_percent": r.get("Coverage_percent"),
            "R2": r.get("R2"),
            "RMSE": r.get("RMSE"),
            "MAE": r.get("MAE"),
            "MAPE_percent": r.get("MAPE_percent"),
            "Physical_Violations": r.get("Physical_Violations"),
            "Interpretation": "In-range RPM interpolation"
        })

    for label, source in [
        ("Step 4 One-step extrapolation", step4_one),
        ("Step 4 Multi-step extrapolation", step4_multi),
    ]:
        for _, r in source.iterrows():
            rows.append({
                "Stage": label,
                "Target": r.get("Target"),
                "Method": r.get("Method"),
                "Coverage_percent": None,
                "R2": r.get("R2"),
                "RMSE": r.get("RMSE"),
                "MAE": r.get("MAE"),
                "MAPE_percent": r.get("MAPE_percent"),
                "Physical_Violations": r.get("Physical_Violations"),
                "Interpretation": "Progressive RPM extrapolation diagnostic"
            })

    for _, r in step5_summary.iterrows():
        rows.append({
            "Stage": "Step 5 independent holdout",
            "Target": r.get("Target"),
            "Method": r.get("Selected_Method"),
            "Coverage_percent": r.get("Coverage_percent"),
            "R2": r.get("R2"),
            "RMSE": r.get("RMSE"),
            "MAE": r.get("MAE"),
            "MAPE_percent": r.get("MAPE_percent"),
            "Physical_Violations": r.get("Physical_Violations"),
            "Interpretation": "17 untouched real experiments"
        })

    return pd.DataFrame(rows)


def make_eda_summary(experimental):
    rows = []

    for target in TARGETS:
        s = pd.to_numeric(
            experimental[target],
            errors="coerce"
        ).dropna()

        rows.append({
            "Target": target,
            "N": len(s),
            "Mean": s.mean(),
            "Std": s.std(ddof=1),
            "Min": s.min(),
            "Max": s.max(),
            "Zero_Count": int((s == 0).sum()),
        })

    return pd.DataFrame(rows)


def make_correlations(experimental):
    return pd.DataFrame([
        {
            "Response_A": "hc_hi",
            "Response_B": "td_to",
            "Pearson_r": experimental["hc_hi"].corr(
                experimental["td_to"],
                method="pearson"
            ),
            "Spearman_rho": experimental["hc_hi"].corr(
                experimental["td_to"],
                method="spearman"
            ),
        },
        {
            "Response_A": "hc_hi",
            "Response_B": "Vch",
            "Pearson_r": experimental["hc_hi"].corr(
                experimental["Vch"],
                method="pearson"
            ),
            "Spearman_rho": experimental["hc_hi"].corr(
                experimental["Vch"],
                method="spearman"
            ),
        },
        {
            "Response_A": "td_to",
            "Response_B": "Vch",
            "Pearson_r": experimental["td_to"].corr(
                experimental["Vch"],
                method="pearson"
            ),
            "Spearman_rho": experimental["td_to"].corr(
                experimental["Vch"],
                method="spearman"
            ),
        },
    ])


def make_graph_sheet_data(experimental, interpolation, target):
    rows = []

    interp_lookup = {
        (float(r["d"]), float(r["Rpm"])): r[target]
        for _, r in interpolation.iterrows()
    }

    for rpm in GRID_RPM:
        row = {"RPM": rpm}

        for d in DIAMETERS:
            exp = experimental[
                (experimental["d"] == d) &
                (experimental["Rpm"] == rpm)
            ]

            if not exp.empty:
                row[f"d={d} Experimental"] = exp.iloc[0][target]
            else:
                row[f"d={d} Experimental"] = None

            row[f"d={d} 5-RPM Model"] = interp_lookup.get(
                (float(d), float(rpm))
            )

        rows.append(row)

    return pd.DataFrame(rows)


def add_line_chart(ws, title, end_col):
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = "Response"
    chart.x_axis.title = "RPM"
    chart.height = 12
    chart.width = 24
    chart.legend.position = "b"

    data = Reference(
        ws,
        min_col=2,
        max_col=end_col,
        min_row=1,
        max_row=ws.max_row
    )

    cats = Reference(
        ws,
        min_col=1,
        min_row=2,
        max_row=ws.max_row
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    ws.add_chart(chart, f"N2")


def add_readme(wb):
    ws = wb.create_sheet("README")

    content = [
        ("STEP 5B — FINAL ANALYSIS / EDA CONSOLIDATION", ""),
        ("Purpose",
         "Single shareable workbook for pre-ML data analysis, graphing, "
         "model-history review, and provenance tracking."),
        ("Experimental dataset",
         "84 real observations: 6 measured diameters × 14 measured RPM levels."),
        ("Train_Anchors",
         "67 real observations used for model development."),
        ("Test_Holdout",
         "17 real observations reserved for independent validation."),
        ("In-range selected models",
         "hc_hi → Cubic Griddata; td_to → 3rd-order RSM; "
         "Vch → 3rd-order RSM."),
        ("Interpolation domain",
         "Only measured diameters 6, 8, 10, 12, 14, 16 mm. "
         "RPM grid is 0–260 at 5-RPM spacing."),
        ("Intermediate diameters",
         "NONE. No 7, 9, 11, 13, or 15 mm points are generated."),
        ("Interpolation points",
         "Graph/analysis representation only by default. They are not treated "
         "as independent experiments."),
        ("RPM extrapolation",
         "Reference/diagnostic only. Not measured data and not ML training data."),
        ("ML default dataset",
         "ML_Train_67."),
        ("ML holdout policy",
         "ML_Holdout_17 remains isolated from model fitting and selection."),
        ("Important target name",
         "Use td_to everywhere; do not rename it to td/to."),
    ]

    for row_idx, (key, value) in enumerate(content, 1):
        ws.cell(row=row_idx, column=1, value=key)
        ws.cell(row=row_idx, column=2, value=value)

    ws.merge_cells("A1:B1")
    ws["A1"].fill = PatternFill("solid", fgColor=TITLE_FILL)
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[0].font = Font(bold=True)
        row[0].fill = PatternFill(
            "solid",
            fgColor="D9EAF7"
        )
        row[0].border = BORDER
        row[1].border = BORDER
        row[1].alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    ws.freeze_panes = "A2"


def main():
    files = [
        STEP1_FILE,
        STEP3_FILE,
        STEP4_FILE,
        STEP5_FILE,
    ]

    missing = [str(p) for p in files if not p.exists()]
    if missing:
        raise FileNotFoundError(
            "Missing required workbook(s):\n" +
            "\n".join(missing)
        )

    print("=" * 78)
    print("STEP 5B - FINAL ANALYSIS DATA CONSOLIDATION")
    print("=" * 78)

    train = read_xlsx(
        STEP1_FILE,
        "Train_Anchors"
    )

    holdout = read_xlsx(
        STEP1_FILE,
        "Test_Holdout"
    )

    step2_summary = read_xlsx(
        STEP1_FILE,
        "Target_Summary"
    )

    step3_summary = read_xlsx(
        STEP3_FILE,
        "Interpolation_CV_Summary"
    )

    step3_grid = read_xlsx(
        STEP3_FILE,
        "RPMOnly_Grid"
    )

    step4_one = read_xlsx(
        STEP4_FILE,
        "OneStep_Summary"
    )

    step4_multi = read_xlsx(
        STEP4_FILE,
        "MultiStep_Summary"
    )

    step4_extrap = read_xlsx(
        STEP4_FILE,
        "RPMOnly_280_320"
    )

    step5_summary = read_xlsx(
        STEP5_FILE,
        "Holdout_Summary"
    )

    step5_predictions = read_xlsx(
        STEP5_FILE,
        "Holdout_Predictions"
    )

    if len(train) != 67:
        raise ValueError(
            f"Expected 67 Train_Anchors rows; found {len(train)}"
        )

    if len(holdout) != 17:
        raise ValueError(
            f"Expected 17 Test_Holdout rows; found {len(holdout)}"
        )

    experimental = make_experimental(
        train,
        holdout
    )

    interpolation = make_rpm_interpolation(
        step3_grid,
        train,
        holdout
    )

    extrapolation = make_extrapolation(
        step4_extrap
    )

    holdout_check = make_holdout_check(
        step5_predictions
    )

    history = make_model_history(
        step2_summary,
        step3_summary,
        step4_one,
        step4_multi,
        step5_summary
    )

    eda = make_eda_summary(experimental)
    correlations = make_correlations(experimental)

    wb = Workbook()

    # Remove the default sheet.
    default = wb.active
    wb.remove(default)

    add_readme(wb)

    write_df(
        wb,
        "Experimental_84",
        experimental
    )

    write_df(
        wb,
        "ML_Train_67",
        train[["Point_ID", "d", "Rpm"] + TARGETS]
    )

    write_df(
        wb,
        "ML_Holdout_17",
        holdout[["Point_ID", "d", "Rpm"] + TARGETS]
    )

    write_df(
        wb,
        "RPM_Interpolation_301",
        interpolation
    )

    write_df(
        wb,
        "RPM_Extrapolation_Ref",
        extrapolation
    )

    write_df(
        wb,
        "Holdout_Model_Check",
        holdout_check
    )

    write_df(
        wb,
        "Model_History",
        history
    )

    write_df(
        wb,
        "EDA_Summary",
        eda
    )

    write_df(
        wb,
        "Response_Correlations",
        correlations
    )

    for target in TARGETS:
        graph_data = make_graph_sheet_data(
            experimental,
            interpolation,
            target
        )

        ws = write_df(
            wb,
            f"Graph_{target}",
            graph_data,
            table=False
        )

        add_line_chart(
            ws,
            f"{target}: Experimental vs 5-RPM In-Range Model",
            len(graph_data.columns)
        )

    # Ensure number formats are sensible.
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"

    wb.save(OUTPUT_FILE)

    print("-" * 78)
    print("OUTPUT")
    print(OUTPUT_FILE)
    print("-" * 78)
    print("Train rows:", len(train))
    print("Holdout rows:", len(holdout))
    print("Experimental rows:", len(experimental))
    print("RPM interpolation rows:", len(interpolation))
    print("Intermediate diameters generated: 0")
    print("-" * 78)
    print("Sheets created:")
    for name in wb.sheetnames:
        print(" -", name)
    print("-" * 78)
    print("ML DEFAULT POLICY")
    print("ML_Train_67       : YES")
    print("ML_Holdout_17     : NO")
    print("RPM_Interpolation : NO by default")
    print("RPM_Extrapolation : NO")
    print("-" * 78)
    print("COMPLETED")


if __name__ == "__main__":
    main()
