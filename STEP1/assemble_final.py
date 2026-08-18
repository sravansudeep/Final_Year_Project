import os
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# FILES
# ============================================================

original_file = "edited_excel.xlsx"
inrange_file = "in_range_griddata_linear.xlsx"
extrapolation_file = "extrapolated_points.xlsx"

original_sheet = "Selected_columns"

output_file = "expanded_air_core_dataset.xlsx"


# ============================================================
# LOAD ORIGINAL DATA
# ============================================================

original = pd.read_excel(
    original_file,
    sheet_name=original_sheet
)

original = original[
    ["d", "rpm", "Vmin"]
].dropna()

original["d"] = original["d"].astype(float)
original["rpm"] = original["rpm"].astype(float)
original["Vmin"] = original["Vmin"].astype(float)


# ============================================================
# LOAD IN-RANGE DATA
# ============================================================

inrange = pd.read_excel(inrange_file)

print("In-range columns:")
print(list(inrange.columns))


inrange = inrange[
    ["d", "rpm", "Vmin", "Data_Type"]
].copy()

inrange["d"] = inrange["d"].astype(float)
inrange["rpm"] = inrange["rpm"].astype(float)
inrange["Vmin"] = inrange["Vmin"].astype(float)


# Add columns ourselves
inrange["Source_Method"] = inrange["Data_Type"].map({
    "Original": "Experimental",
    "Interpolated": "Linear Griddata"
})

inrange["Validation_Status"] = inrange["Data_Type"].map({
    "Original": "Measured",
    "Interpolated": "In-range model estimate"
})


# ============================================================
# LOAD EXTRAPOLATION DATA
# ============================================================

extrapolated = pd.read_excel(
    extrapolation_file
)

print("Extrapolation columns:")
print(list(extrapolated.columns))


extrapolated = extrapolated[
    ["d", "rpm", "Vmin", "Data_Type"]
].copy()

extrapolated["d"] = extrapolated["d"].astype(float)
extrapolated["rpm"] = extrapolated["rpm"].astype(float)
extrapolated["Vmin"] = extrapolated["Vmin"].astype(float)


extrapolated["Source_Method"] = "Third-order RSM"
extrapolated["Validation_Status"] = "Provisional"


# ============================================================
# PROTECT ORIGINAL EXPERIMENTAL VALUES
# ============================================================

for _, row in original.iterrows():

    mask = (
        (inrange["d"] == row["d"]) &
        (inrange["rpm"] == row["rpm"])
    )

    inrange.loc[mask, "Vmin"] = row["Vmin"]
    inrange.loc[mask, "Data_Type"] = "Original"
    inrange.loc[mask, "Source_Method"] = "Experimental"
    inrange.loc[mask, "Validation_Status"] = "Measured"


# ============================================================
# COMBINE
# ============================================================

combined = pd.concat(
    [
        inrange,
        extrapolated
    ],
    ignore_index=True
)


# ============================================================
# SORT
# ============================================================

data_type_order = {
    "Original": 0,
    "Interpolated": 1,
    "Extrapolated": 2
}

combined["_sort"] = combined["Data_Type"].map(
    data_type_order
)

combined = combined.sort_values(
    by=["_sort", "d", "rpm"]
).drop(
    columns="_sort"
).reset_index(drop=True)


# ============================================================
# MODEL VALIDATION
# ============================================================

validation = pd.DataFrame([
    [
        "Polynomial Degree 1",
        0.827562,
        0.039062,
        0.030892,
        0.001526,
        25.62886,
        80,
        84
    ],
    [
        "Polynomial Degree 2",
        0.971219,
        0.015959,
        0.013047,
        0.000255,
        9.915637,
        80,
        84
    ],
    [
        "Polynomial Degree 3",
        0.976989,
        0.014270,
        0.011534,
        0.000204,
        10.940410,
        80,
        84
    ],
    [
        "Griddata Linear",
        0.975292,
        0.014787,
        0.010007,
        0.000219,
        8.223970,
        80,
        84
    ],
    [
        "Griddata Cubic",
        0.972543,
        0.015587,
        0.012069,
        0.000243,
        9.255467,
        80,
        84
    ]
], columns=[
    "Method",
    "R2",
    "RMSE",
    "MAE",
    "MSE",
    "MAPE (%)",
    "Valid Predictions",
    "Total Locations"
])


# ============================================================
# RSM MODEL
# ============================================================

rsm = pd.DataFrame([
    ["Model", "Third-order Response Surface Model"],
    ["Training observations", 84],
    ["Full-data R2", 0.9817227521694277],
    ["Full-data RMSE", 0.013317909725761863],
    ["Full-data MAE", 0.010786927344061443],
    ["Full-data MSE", 0.00017736671946354244],
    ["Full-data MAPE (%)", 9.709419569990322],
    ["Intercept", 0.027745884636646206],
    ["d", -3.61906403486e-05],
    ["rpm", 0.000401359595947],
    ["d^2", -0.000384986452466],
    ["d*rpm", 0.000364704809031],
    ["rpm^2", -9.69271240385e-06],
    ["d^3", 1.23457981408e-05],
    ["d^2*rpm", -1.03624336833e-05],
    ["d*rpm^2", -2.37570088046e-07],
    ["rpm^3", 1.94801505762e-08]
], columns=[
    "Parameter",
    "Value"
])


equation = (
    "A·Hc = 0.02774588464 "
    "- 0.00003619064*d "
    "+ 0.00040135960*rpm "
    "- 0.00038498645*d^2 "
    "+ 0.00036470481*d*rpm "
    "- 0.00000969271*rpm^2 "
    "+ 0.00001234580*d^3 "
    "- 0.00001036243*d^2*rpm "
    "- 0.00000023757*d*rpm^2 "
    "+ 0.0000000194802*rpm^3"
)

rsm.loc[len(rsm)] = [
    "Equation",
    equation
]


# ============================================================
# README
# ============================================================

readme = pd.DataFrame([
    [
        "Project",
        "AI-Assisted Framework for Prediction, Characterization, and Severity Assessment of Air-Core Vortices in Draining Tanks"
    ],
    ["Original experimental observations", 84],
    ["Experimental diameters", "6, 8, 10, 12, 14, 16 mm"],
    ["Experimental RPM", "0–260 RPM in 20 RPM increments"],
    ["In-range diameter grid", "6–16 mm in 1 mm increments"],
    ["In-range RPM grid", "0–260 RPM in 10 RPM increments"],
    ["Total in-range locations", 297],
    ["Original locations retained", 84],
    ["New in-range points", 213],
    ["Extrapolated points", 6],
    ["Final expanded rows", 303],
    ["Model-selection method", "Leave-One-Out Cross-Validation (LOOCV)"],
    ["Best overall statistical model", "Third-order Polynomial RSM"],
    ["LOOCV R2", 0.976989],
    ["LOOCV RMSE", 0.014270],
    ["In-range interpolation method", "Linear Griddata"],
    ["In-range validation", "297 grid points; 213 generated; 0 NaN; 0 negative"],
    ["Extrapolation method", "Third-order RSM"],
    ["Extrapolation diameters", "4 and 18 mm"],
    ["Extrapolation RPM", "280, 300, 320 RPM"],
    ["Extrapolation status", "Provisional"],
    [
        "Important limitation",
        "Generated points are model estimates, not experimental measurements. Extrapolated points require experimental or CFD validation."
    ]
], columns=[
    "Item",
    "Information"
])


# ============================================================
# WRITE EXCEL
# ============================================================

with pd.ExcelWriter(
    output_file,
    engine="openpyxl"
) as writer:

    readme.to_excel(
        writer,
        sheet_name="README",
        index=False
    )

    original.to_excel(
        writer,
        sheet_name="Original_Raw",
        index=False
    )

    inrange.to_excel(
        writer,
        sheet_name="In_Range_Expanded",
        index=False
    )

    extrapolated.to_excel(
        writer,
        sheet_name="Extrapolated",
        index=False
    )

    combined.to_excel(
        writer,
        sheet_name="Combined_Expanded",
        index=False
    )

    validation.to_excel(
        writer,
        sheet_name="Model_Validation",
        index=False
    )

    rsm.to_excel(
        writer,
        sheet_name="RSM_Model",
        index=False
    )


# ============================================================
# EXCEL FORMATTING
# ============================================================

wb = load_workbook(output_file)


green_fill = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE"
)

yellow_fill = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC"
)

orange_fill = PatternFill(
    fill_type="solid",
    fgColor="F4B183"
)

header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

header_font = Font(
    bold=True
)


for ws in wb.worksheets:

    ws.freeze_panes = "A2"

    for cell in ws[1]:

        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for column in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            try:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )
            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            60
        )


# ============================================================
# COLOR COMBINED DATA
# ============================================================

ws = wb["Combined_Expanded"]

headers = {
    cell.value: cell.column
    for cell in ws[1]
}

data_type_col = headers["Data_Type"]


for row in range(2, ws.max_row + 1):

    data_type = ws.cell(
        row=row,
        column=data_type_col
    ).value

    if data_type == "Original":
        fill = green_fill

    elif data_type == "Interpolated":
        fill = yellow_fill

    elif data_type == "Extrapolated":
        fill = orange_fill

    else:
        continue

    for col in range(
        1,
        ws.max_column + 1
    ):
        ws.cell(
            row=row,
            column=col
        ).fill = fill


# ============================================================
# COLOR IN-RANGE DATA
# ============================================================

ws = wb["In_Range_Expanded"]

headers = {
    cell.value: cell.column
    for cell in ws[1]
}

data_type_col = headers["Data_Type"]


for row in range(2, ws.max_row + 1):

    data_type = ws.cell(
        row=row,
        column=data_type_col
    ).value

    fill = (
        green_fill
        if data_type == "Original"
        else yellow_fill
    )

    for col in range(
        1,
        ws.max_column + 1
    ):
        ws.cell(
            row=row,
            column=col
        ).fill = fill


# ============================================================
# COLOR EXTRAPOLATED
# ============================================================

ws = wb["Extrapolated"]

for row in range(2, ws.max_row + 1):

    for col in range(
        1,
        ws.max_column + 1
    ):
        ws.cell(
            row=row,
            column=col
        ).fill = orange_fill


# ============================================================
# HIGHLIGHT SELECTED METHODS
# ============================================================

ws = wb["Model_Validation"]

for row in range(2, ws.max_row + 1):

    method = ws.cell(
        row=row,
        column=1
    ).value

    if method == "Polynomial Degree 3":

        for col in range(
            1,
            ws.max_column + 1
        ):
            ws.cell(
                row=row,
                column=col
            ).font = Font(
                bold=True
            )

    elif method == "Griddata Linear":

        for col in range(
            1,
            ws.max_column + 1
        ):
            ws.cell(
                row=row,
                column=col
            ).fill = yellow_fill


# ============================================================
# README FORMATTING
# ============================================================

ws = wb["README"]

ws.column_dimensions["A"].width = 40
ws.column_dimensions["B"].width = 90

for row in ws.iter_rows():

    for cell in row:

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )


# ============================================================
# SAVE
# ============================================================

wb.save(output_file)


# ============================================================
# FINAL CHECKS
# ============================================================

original_count = len(original)

inrange_total = len(inrange)

inrange_original = (
    inrange["Data_Type"] == "Original"
).sum()

inrange_interpolated = (
    inrange["Data_Type"] == "Interpolated"
).sum()

extrapolated_count = len(extrapolated)

combined_count = len(combined)

growth = (
    (combined_count - original_count)
    / original_count
) * 100


print()
print("=" * 80)
print("DATASET ASSEMBLY COMPLETE")
print("=" * 80)

print()
print("Original observations :", original_count)
print("In-range locations    :", inrange_total)
print("Original in-range     :", inrange_original)
print("Interpolated points   :", inrange_interpolated)
print("Extrapolated points   :", extrapolated_count)
print("Final combined rows   :", combined_count)

print()
print("Expected:")
print("Original       : 84")
print("In-range total : 297")
print("Interpolated   : 213")
print("Extrapolated   : 6")
print("Final total    : 303")

print()
print(f"Dataset growth: {growth:.2f}%")

print()
print("Output:")
print(os.path.abspath(output_file))

print()
print("=" * 80)